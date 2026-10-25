import sqlite3
from pathlib import Path
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "database" / "reconciliation.db"
CHECKS_SQL_PATH = BASE_DIR / "database" / "reconciliation_checks.sql"
OUTPUT_DIR = BASE_DIR / "data" / "output"
REPORT_PATH = OUTPUT_DIR / "Finance_Audit_Report.xlsx"

CURRENCY_COLUMNS = {
    "total_budget_pln",
    "actual_spend_pln",
    "remaining_budget_pln",
    "monthly_limit_pln",
    "total_spent_pln",
    "overrun_amount_pln",
    "amount_pln",
}

PERCENT_COLUMNS = {
    "budget_utilization_pct",
    "limit_utilization_pct",
}


def load_sql_queries(sql_path: Path) -> dict[str, str]:
    """Load SQL queries separated by semicolons and map them to report sections."""
    content = sql_path.read_text(encoding="utf-8")
    raw_statements = [stmt.strip() for stmt in content.split(";") if stmt.strip()]

    query_names = [
        "EXECUTIVE_SUMMARY",
        "BUDGET_OVERRUNS",
        "SECURITY_ALERTS",
        "DUPLICATE_SUSPECTS",
        "COST_CENTER_MISMATCHES",
        "REJECTED_ROWS",
    ]

    return dict(zip(query_names, raw_statements))


def style_worksheet(ws, start_row: int = 1) -> None:
    """Apply header styling, number formatting, and auto-fit column widths."""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[start_row]:
        if cell.value is not None:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = [cell.value for cell in ws[start_row]]
    for col_idx, col_name in enumerate(headers, start=1):
        if col_name in CURRENCY_COLUMNS:
            for row in range(start_row + 1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00 "zł"'
        elif col_name in PERCENT_COLUMNS:
            for row in range(start_row + 1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00"%"'

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_audit_report() -> None:
    """Execute SQL reconciliation checks and export results to a formatted Excel workbook."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    queries = load_sql_queries(CHECKS_SQL_PATH)

    with sqlite3.connect(DB_PATH) as conn:
        df_exec = pd.read_sql_query(queries["EXECUTIVE_SUMMARY"], conn)
        df_overruns = pd.read_sql_query(queries["BUDGET_OVERRUNS"], conn)
        df_security = pd.read_sql_query(queries["SECURITY_ALERTS"], conn)
        df_duplicates = pd.read_sql_query(queries["DUPLICATE_SUSPECTS"], conn)
        df_mismatches = pd.read_sql_query(queries["COST_CENTER_MISMATCHES"], conn)
        df_rejected = pd.read_sql_query(queries["REJECTED_ROWS"], conn)

    with pd.ExcelWriter(REPORT_PATH, engine="openpyxl") as writer:
        df_exec.to_excel(writer, sheet_name="Executive_Summary", index=False)
        df_overruns.to_excel(writer, sheet_name="Budget_Overruns", index=False)
        df_security.to_excel(writer, sheet_name="Security_Alerts", index=False)
        df_duplicates.to_excel(writer, sheet_name="Duplicate_Suspects", index=False)

        df_mismatches.to_excel(writer, sheet_name="Data_Mismatches", index=False, startrow=0)
        rejected_start_row = len(df_mismatches) + 3
        df_rejected.to_excel(
            writer,
            sheet_name="Data_Mismatches",
            index=False,
            startrow=rejected_start_row,
        )

        workbook = writer.book
        for sheet_name in ["Executive_Summary", "Budget_Overruns", "Security_Alerts", "Duplicate_Suspects"]:
            style_worksheet(workbook[sheet_name], start_row=1)

        ws_mismatch = workbook["Data_Mismatches"]
        style_worksheet(ws_mismatch, start_row=1)
        style_worksheet(ws_mismatch, start_row=rejected_start_row + 1)

    print(f"[OK] Generated audit report: {REPORT_PATH}")
    print(f"     - Executive_Summary:  {len(df_exec)} departments")
    print(f"     - Budget_Overruns:    {len(df_overruns)} limit overruns")
    print(f"     - Security_Alerts:    {len(df_security)} post-termination expenses")
    print(f"     - Duplicate_Suspects: {len(df_duplicates)} suspected duplicate rows")
    print(f"     - Data_Mismatches:    {len(df_mismatches)} CC mismatches + {len(df_rejected)} ETL rejected rows")


def main() -> None:
    print("--- Starting SQL Reconciliation Checks & Reporting (Modules 4 & 5) ---")
    generate_audit_report()
    print("--- Audit Report Generation Completed Successfully ---")


if __name__ == "__main__":
    main()